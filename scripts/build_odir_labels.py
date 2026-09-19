"""Build the per-eye ODIR-5K label file (brief Task 6).

Reads Internal/ODIR-5K/data.xlsx (never edited) and writes Labels/odir_eye_labels.csv
with one row per eye. Every label comes from that eye's own diagnosis keywords, never
from the patient-level N, D, G, C, A, H, M, O columns.

The split column is left empty: Task 7 fills it only after the Task 4 decisions and
the ODIR-5K licence check (Gate A).

Run from the Datasets folder:  python Labels/build_odir_labels.py
Needs: openpyxl
"""
import csv
import os
import sys
from collections import defaultdict

from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(ROOT, "Internal", "ODIR-5K", "data.xlsx")
IMAGES = os.path.join(ROOT, "Internal", "ODIR-5K", "Training Images")
OUT = os.path.join(ROOT, "Labels", "odir_eye_labels.csv")

EXCLUDE = {"no fundus image", "anterior segment image", "image offset",
           "optic disk photographically invisible", "low image quality"}
QUALITY = {"lens dust", "low image quality"}
DR = {"mild nonproliferative retinopathy", "moderate non proliferative retinopathy",
      "severe nonproliferative retinopathy", "proliferative diabetic retinopathy",
      "severe proliferative diabetic retinopathy", "diabetic retinopathy"}

# Expected per-eye counts; the script stops if any differ.
EXPECTED = {"rows": 7000, "cataract": 313, "glaucoma": 282, "suspected_glaucoma": 44,
            "media_opacity": 63, "quality_flag": 426, "exclude": 30,
            "normal_only": 2818, "diabetic_retinopathy": 1796}


def tokens(cell):
    """Fullwidth comma to comma, lowercase, split, strip, collapse spaces; whole tokens only."""
    text = (cell or "").replace("，", ",").lower()
    return {" ".join(t.split()) for t in text.split(",") if t.strip()}


def main():
    ws = load_workbook(XLSX, read_only=True).active
    rows = ws.iter_rows(values_only=True)
    header = [str(h).strip() for h in next(rows)]
    col = {name: i for i, name in enumerate(header)}

    out, counts = [], defaultdict(int)
    per_patient = defaultdict(lambda: {"cataract": 0, "glaucoma": 0})
    for r in rows:
        if r[col["ID"]] is None:
            continue
        pid = int(r[col["ID"]])
        for eye, fcol, kcol in (("left", "Left-Fundus", "Left-Diagnostic Keywords"),
                                ("right", "Right-Fundus", "Right-Diagnostic Keywords")):
            raw = r[col[kcol]] or ""
            t = tokens(raw)
            row = {
                "patient_id": pid,
                "eye": eye,
                "filename": r[col[fcol]],
                "diagnosis_text": raw,
                "cataract": int("cataract" in t),
                "glaucoma": int("glaucoma" in t),
                "suspected_glaucoma": int("suspected glaucoma" in t),
                "media_opacity": int("refractive media opacity" in t),
                "quality_flag": int(bool(t & QUALITY)),
                "exclude": int(bool(t & EXCLUDE)),
                "split": "",
            }
            out.append(row)
            for k in ("cataract", "glaucoma", "suspected_glaucoma", "media_opacity", "quality_flag", "exclude"):
                counts[k] += row[k]
            counts["normal_only"] += int(t == {"normal fundus"})
            counts["diabetic_retinopathy"] += int(bool(t & DR))
            per_patient[pid]["cataract"] += row["cataract"]
            per_patient[pid]["glaucoma"] += row["glaucoma"]
    counts["rows"] = len(out)

    problems = [f"{k}: expected {v}, got {counts[k]}" for k, v in EXPECTED.items() if counts[k] != v]
    cat = [p for p in per_patient.values() if p["cataract"]]
    gla = [p for p in per_patient.values() if p["glaucoma"]]
    checks = {"cataract patients": (len(cat), 212), "cataract one-eye patients": (sum(p["cataract"] == 1 for p in cat), 111),
              "glaucoma patients": (len(gla), 177), "glaucoma one-eye patients": (sum(p["glaucoma"] == 1 for p in gla), 72),
              "patients with both conditions": (sum(bool(p["cataract"] and p["glaucoma"]) for p in per_patient.values()), 1)}
    problems += [f"{k}: expected {e}, got {g}" for k, (g, e) in checks.items() if g != e]
    on_disk = set(os.listdir(IMAGES))
    missing = [row["filename"] for row in out if row["filename"] not in on_disk]
    if missing:
        problems.append(f"{len(missing)} filenames not in Training Images, e.g. {missing[:3]}")
    if problems:
        sys.exit("Label file NOT written:\n  " + "\n  ".join(problems))

    out.sort(key=lambda row: (row["patient_id"], row["eye"] != "left"))
    with open(OUT, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(out[0].keys()))
        w.writeheader()
        w.writerows(out)
    print(f"Wrote {OUT}")
    for k in EXPECTED:
        print(f"  {k}: {counts[k]}")
    for k, (g, _) in checks.items():
        print(f"  {k}: {g}")
    print("  all filenames found in Training Images")


if __name__ == "__main__":
    main()
